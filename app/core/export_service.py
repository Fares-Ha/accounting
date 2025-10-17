import openpyxl
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from PyQt6.QtWidgets import QFileDialog
from datetime import datetime
from ..database import models
import os

def generate_sales_receipt_pdf(order: models.SalesOrder):
    """
    Generates a PDF receipt for a given sales order.
    """
    # Create a 'receipts' directory if it doesn't exist
    if not os.path.exists('receipts'):
        os.makedirs('receipts')

    filepath = f"receipts/receipt_order_{order.id}.pdf"
    doc = SimpleDocTemplate(filepath, pagesize=letter)
    styles = getSampleStyleSheet()
    story = []

    # Title
    story.append(Paragraph("Sales Receipt", styles['Title']))
    story.append(Spacer(1, 12))

    # Order Info
    story.append(Paragraph(f"<b>Order ID:</b> {order.id}", styles['Normal']))
    story.append(Paragraph(f"<b>Customer:</b> {order.customer.name}", styles['Normal']))
    story.append(Paragraph(f"<b>Date:</b> {order.created_at.strftime('%Y-%m-%d %H:%M:%S')}", styles['Normal']))
    story.append(Spacer(1, 12))

    # Items Table
    data = [["Product", "Quantity", "Unit Price", "Total"]]
    for item in order.items:
        total_price = item.quantity * item.price_per_unit
        data.append([
            item.product.name,
            str(item.quantity),
            f"{item.price_per_unit:.2f}",
            f"{total_price:.2f}"
        ])

    # Add total row
    data.append(["", "", "<b>Total Amount</b>", f"<b>{order.total_amount:.2f}</b>"])

    table = Table(data)
    style = TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -2), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('BACKGROUND', (0, -1), (-1, -1), colors.lightgrey),
        ('TEXTCOLOR', (0, -1), (-1, -1), colors.black),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
    ])
    table.setStyle(style)
    story.append(table)

    doc.build(story)
    return os.path.abspath(filepath)


class ExportService:
    """
    Service for exporting data to different formats.
    """
    def export_pnl_report_to_pdf(self, report_data, start_date, end_date, file_path):
        """
        Exports a Profit & Loss report to a PDF file.
        """
        doc = SimpleDocTemplate(file_path, pagesize=letter)
        styles = getSampleStyleSheet()
        story = []

        story.append(Paragraph("Profit & Loss Statement", styles['Title']))
        story.append(Spacer(1, 12))
        story.append(Paragraph(f"For the period from {start_date.strftime('%Y-%m-%d')} to {end_date.strftime('%Y-%m-%d')}", styles['Normal']))
        story.append(Spacer(1, 24))

        data = [
            ["Description", "Amount"],
            ["Total Revenue", f"{report_data['revenue']:.2f}"],
            ["Total Expenses", f"{report_data['expenses']:.2f}"],
            ["<b>Net Profit</b>", f"<b>{report_data['net_profit']:.2f}</b>"],
        ]

        table = Table(data, colWidths=[300, 100])
        style = TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ])
        table.setStyle(style)
        story.append(table)

        doc.build(story)

    def export_pnl_report_to_excel(self, report_data, start_date, end_date, file_path):
        """
        Exports a Profit & Loss report to an Excel file.
        """
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Profit and Loss"

        sheet.append([f"Profit & Loss Statement"])
        sheet.append([f"For the period from {start_date.strftime('%Y-%m-%d')} to {end_date.strftime('%Y-%m-%d')}"])
        sheet.append([])

        headers = ["Description", "Amount"]
        sheet.append(headers)

        data = [
            ["Total Revenue", report_data['revenue']],
            ["Total Expenses", report_data['expenses']],
            ["Net Profit", report_data['net_profit']],
        ]

        for row_data in data:
            sheet.append(row_data)

        for i in range(1, 5):
            sheet.cell(row=i, column=2).number_format = '"$"#,##0.00'

        workbook.save(file_path)

    def export_balance_sheet_report_to_pdf(self, report_data, as_of_date, file_path):
        """
        Exports a Balance Sheet report to a PDF file.
        """
        doc = SimpleDocTemplate(file_path, pagesize=letter)
        styles = getSampleStyleSheet()
        story = []

        story.append(Paragraph("Balance Sheet", styles['Title']))
        story.append(Spacer(1, 12))
        story.append(Paragraph(f"As of {as_of_date.strftime('%Y-%m-%d')}", styles['Normal']))
        story.append(Spacer(1, 24))

        data = [
            ["Description", "Amount"],
            ["Total Assets", f"{report_data['assets']:.2f}"],
            ["Total Liabilities", f"{report_data['liabilities']:.2f}"],
            ["<b>Total Equity</b>", f"<b>{report_data['equity']:.2f}</b>"],
        ]

        table = Table(data, colWidths=[300, 100])
        style = TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ])
        table.setStyle(style)
        story.append(table)

        doc.build(story)

    def export_sales_report_to_pdf(self, sales_data, file_path):
        headers = ["Product", "Total Quantity", "Total Revenue"]
        data = [[item.name, str(item.total_quantity), f"{item.total_revenue / 100:.2f}"] for item in sales_data]
        self.export_to_pdf(data, headers, file_path)

    def export_sales_report_to_excel(self, sales_data, file_path):
        headers = ["Product", "Total Quantity", "Total Revenue"]
        data = [[item.name, item.total_quantity, item.total_revenue / 100] for item in sales_data]
        self.export_to_excel(data, headers, file_path)

    def export_to_excel(self, data, headers, file_path):
        """
        Exports data to an Excel file.
        """
        if not data:
            return

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(headers)

        for row_data in data:
            sheet.append(row_data)

        workbook.save(file_path)

    def export_to_pdf(self, data, headers, file_path):
        """
        Exports data to a PDF file.
        """
        if not data:
            return

        doc = SimpleDocTemplate(file_path, pagesize=letter)
        table_data = [headers] + data

        table = Table(table_data)
        style = TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ])
        table.setStyle(style)

        doc.build([table])